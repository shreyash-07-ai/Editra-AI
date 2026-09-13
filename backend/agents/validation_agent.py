from pathlib import Path
from docx import Document
from pptx import Presentation
import openpyxl

class ValidationAgent:
    def validate(self, path, artifact_type):
        p = Path(path)
        if not p.exists() or p.stat().st_size == 0:
            return False, ["File missing or empty."]
        try:
            if artifact_type == "docx":
                Document(str(p))
            elif artifact_type == "pptx":
                Presentation(str(p))
            elif artifact_type == "xlsx":
                wb = openpyxl.load_workbook(str(p))
                if not wb.sheetnames:
                    return False, ["Workbook has no sheets."]
            elif artifact_type == "csv":
                with open(p, newline="", encoding="utf-8-sig", errors="replace") as f:
                    if not f.readline():
                        return False, ["CSV file is empty."]
            elif artifact_type == "txt":
                text = p.read_text(encoding="utf-8", errors="replace")
                if not text.strip():
                    return False, ["TXT file is empty."]
            elif artifact_type == "pdf":
                import fitz
                doc = fitz.open(str(p))
                try:
                    if doc.page_count < 1:
                        return False, ["PDF has no pages."]
                finally:
                    doc.close()
            else:
                return False, ["Unsupported artifact type."]
            return True, ["Artifact opens successfully."]
        except Exception as e:
            return False, [f"Artifact validation failed: {e}"]
