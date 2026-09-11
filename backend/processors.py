from pathlib import Path
from docx import Document
from pptx import Presentation
import fitz

def extract_docx(path):
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        tables.append([[c.text for c in row.cells] for row in t.rows])
    style = {}
    if doc.styles:
        try:
            style["normal_font"] = doc.styles["Normal"].font.name
        except Exception:
            pass
    return {"type":"docx","paragraphs":paragraphs,"tables":tables,"style":style}

def extract_pptx(path):
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        items = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                items.append({"text":shape.text, "left":shape.left, "top":shape.top,
                              "width":shape.width, "height":shape.height})
        slides.append({"number":i, "items":items, "layout":slide.slide_layout.name if slide.slide_layout else ""})
    return {"type":"pptx","slide_count":len(prs.slides),"slides":slides}

def extract_pdf(path):
    doc = fitz.open(path)
    pages = [p.get_text("text") for p in doc]
    return {"type":"pdf","page_count":len(pages),"pages":pages}

def extract_image(path):
    return {"type":"image","path":str(path),"note":"Image requires OCR/vision processing if OCR dependencies/API are enabled."}

def analyze(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx": return extract_docx(path)
    if ext == ".pptx": return extract_pptx(path)
    if ext == ".pdf": return extract_pdf(path)
    if ext in {".png",".jpg",".jpeg"}: return extract_image(path)
    raise ValueError(f"Unsupported file: {path}")
