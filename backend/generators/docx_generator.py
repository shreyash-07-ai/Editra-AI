from pathlib import Path
import json
import re
import shutil
import tempfile

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from ..llm import LLM


class DOCXGenerator:
    """Edit an existing DOCX in-place on a copied version of the current artifact."""
    BACKUP_PREFIX = "EDITRA_TABLE_BACKUP::"
    CHART_CAPTION_PREFIX = "EDITRA_CHART::"

    def __init__(self):
        self.llm = LLM()

    def generate(self, output_path, prompt, source_text="", existing=None, template_path=None, edit_plan=None):
        base = existing.get("path") if existing else template_path
        if base and Path(base).exists():
            shutil.copy2(base, output_path)
            if edit_plan:
                self.apply_edit_plan(output_path, edit_plan, source_text, prompt)
            else:
                doc = Document(output_path)
                addition = self._generate_addition(prompt, source_text)
                if addition:
                    self._insert_generated_section(doc, addition)
                self._polish_layout(doc)
                doc.save(output_path)
            return output_path

        content = self.llm.json(
            "Create a professional document from the source. Return JSON only with title and sections.",
            f"REQUEST:\n{prompt}\nSOURCE:\n{source_text}",
            {"title": "Editra AI Document", "sections": [{"heading": "Source Content", "paragraphs": [source_text[:5000]]}]},
        )
        doc = Document()
        doc.add_heading(content.get("title", "Editra AI Document"), 0)
        for item in content.get("sections", []):
            doc.add_heading(item.get("heading", "Section"), 1)
            for text in item.get("paragraphs", []):
                doc.add_paragraph(text)
        self._polish_layout(doc)
        doc.save(output_path)
        return output_path

    def _generate_addition(self, prompt, source_text):
        return self.llm.json(
            "Generate only the requested addition to an existing document. Return JSON: {\"heading\":\"...\",\"paragraphs\":[\"...\"]}. Preserve existing content.",
            f"REQUEST:\n{prompt}\nSOURCE:\n{source_text}",
            None,
        )

    def _insert_generated_section(self, doc, addition):
        if not addition:
            return
        doc.add_heading(addition.get("heading", "Editra AI Update"), 1)
        for text in addition.get("paragraphs", []):
            doc.add_paragraph(text)

    def apply_edit_plan(self, path, plan, source_text="", prompt=""):
        doc = Document(path)
        operations = plan.get("operations", []) if isinstance(plan, dict) else []
        changed = False

        for op in operations:
            kind = op.get("type")
            if kind == "noop":
                continue
            if kind == "replace_paragraph":
                changed |= self._replace_paragraph(doc, int(op.get("index", -1)), op.get("text", ""))
            elif kind == "delete_paragraph":
                changed |= self._delete_paragraph(doc, int(op.get("index", -1)))
            elif kind == "insert_after_heading":
                changed |= self._insert_after_heading(doc, op.get("heading", ""), op.get("paragraphs", []))
            elif kind == "insert_section_after_heading":
                changed |= self._insert_section_after_heading(doc, op.get("after_heading", ""), op.get("heading", "Section"), op.get("paragraphs", []))
            elif kind == "insert_description_at_beginning":
                changed |= self._insert_description_at_beginning(doc)
            elif kind == "insert_executive_summary_at_beginning":
                changed |= self._insert_executive_summary(doc, prompt, source_text)
            elif kind == "shorten_section_paragraph":
                changed |= self._shorten_section(doc, op.get("heading", ""), prompt)
            elif kind == "detail_implementation_table":
                changed |= self._detail_implementation_table(doc, int(op.get("table_index", 0)), prompt)
            elif kind == "append_section":
                doc.add_heading(op.get("heading", "Section"), 1)
                for text in op.get("paragraphs", []):
                    doc.add_paragraph(text)
                changed = True
            elif kind == "replace_table":
                changed |= self._replace_table(doc, int(op.get("table_index", -1)), op.get("rows", []))
            elif kind == "add_chart_from_tables":
                changed |= self._add_chart_from_table(doc, int(op.get("table_index", -1)), op.get("chart_kind", "bar"), op.get("title", "Chart"), False)
            elif kind == "replace_table_with_chart":
                changed |= self._add_chart_from_table(doc, int(op.get("table_index", -1)), op.get("chart_kind", "bar"), op.get("title", "Chart"), True)
            elif kind == "replace_visual_with_table":
                changed |= self._restore_table_from_backup(doc)
            elif kind == "remove_visuals":
                changed |= self._remove_generated_visuals(doc)

        if not changed and operations and not all(op.get("type") == "noop" for op in operations):
            raise ValueError("The requested document operation could not be applied safely.")

        self._polish_layout(doc)
        doc.save(path)
        return path

    def _insert_paragraph_before(self, doc, anchor, text, style="Normal"):
        # Use python-docx's public document API to create the paragraph so its
        # parent is a valid _Body object with a .part property. Then move the
        # already-created XML element before the target paragraph.
        p = doc.add_paragraph(text, style=style)
        anchor.addprevious(p._p)
        return p

    def _find_heading_paragraph(self, doc, heading):
        target = (heading or "").strip().lower()
        if not target:
            return None
        for p in doc.paragraphs:
            if p.text.strip().lower() == target:
                return p
        # Loosen to a substring match against actual headings so the agent
        # isn't limited to one exact demo document's heading wording.
        for p in doc.paragraphs:
            if p.style and p.style.name.startswith("Heading"):
                ptext = p.text.strip().lower()
                if ptext and (target in ptext or ptext in target):
                    return p
        return None

    def _insert_description_at_beginning(self, doc):
        title = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), "the uploaded document")
        headings = [p.text.strip() for p in doc.paragraphs if p.text.strip() and p.style and p.style.name.startswith("Heading")]
        summary = f"This document, titled '{title}', presents the proposal and supporting information contained in the uploaded artifact."
        if headings:
            summary += " It covers " + ", ".join(headings[:8]) + "."
        summary += " The existing document content is preserved below this description."

        first = next((p for p in doc.paragraphs if p.text.strip()), None)
        if first is None:
            return False
        anchor = first._p
        heading = self._insert_paragraph_before(doc, anchor, "Document Description", "Heading 1")
        self._insert_paragraph_before(doc, anchor, summary, "Normal")
        return bool(heading)

    def _insert_executive_summary(self, doc, prompt="", source_text=""):
        first = next((p for p in doc.paragraphs if p.text.strip()), None)
        if first is None:
            return False
        anchor = first._p
        text = self._generate_executive_summary(doc, prompt, source_text)
        heading = self._insert_paragraph_before(doc, anchor, "Executive Summary", "Heading 1")
        self._insert_paragraph_before(doc, anchor, text, "Normal")
        return bool(heading)

    def _generate_executive_summary(self, doc, prompt, source_text):
        result = self.llm.json(
            "Write a concise 3-5 sentence executive summary of the document content below, "
            "grounded only in what it actually contains — do not invent a company name, product, "
            'or facts that are not present. Return JSON only: {"summary":"..."}.',
            f"REQUEST:\n{prompt}\nDOCUMENT CONTENT:\n{source_text[:6000]}",
            None,
        )
        if isinstance(result, dict) and result.get("summary"):
            return result["summary"]
        # Generic, content-derived fallback for when the LLM is unavailable —
        # never invents company names or unrelated facts.
        title = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), "This document")
        headings = [p.text.strip() for p in doc.paragraphs if p.text.strip() and p.style and p.style.name.startswith("Heading")]
        summary = f"This document, '{title}', "
        summary += ("covers " + ", ".join(headings[:6]) + ".") if headings else "presents the content included below."
        return summary

    def _insert_section_after_heading(self, doc, after_heading, heading, paragraphs):
        anchor_p = self._find_heading_paragraph(doc, after_heading)
        if anchor_p is None:
            # Heading not found in this document: append the requested section
            # at the end instead of silently failing the whole request.
            doc.add_heading(heading or "Section", 1)
            for text in paragraphs:
                doc.add_paragraph(text)
            return True
        anchor = anchor_p._p
        body = doc._body._element
        siblings = list(body)
        idx = siblings.index(anchor)
        insert_anchor = anchor
        for sibling in siblings[idx + 1:]:
            if sibling.tag.endswith("}p"):
                text = "".join(sibling.itertext()).strip()
                if text and any(text.startswith(f"{n}.") for n in range(1, 30)):
                    break
            insert_anchor = sibling
        h = doc.add_heading(heading, 1)
        insert_anchor.addnext(h._p)
        anchor2 = h._p
        for text in paragraphs:
            np = doc.add_paragraph(text)
            anchor2.addnext(np._p)
            anchor2 = np._p
        return True

    def _find_heading_index(self, paragraphs, heading):
        target = (heading or "").strip().lower()
        if not target:
            return None
        for i, p in enumerate(paragraphs):
            if p.text.strip().lower() == target:
                return i
        for i, p in enumerate(paragraphs):
            if p.style and p.style.name.startswith("Heading"):
                ptext = p.text.strip().lower()
                if ptext and (target in ptext or ptext in target):
                    return i
        return None

    def _shorten_section(self, doc, heading, prompt=""):
        # Look up the heading and iterate within a single doc.paragraphs call:
        # python-docx builds a fresh list of wrapper objects on every access,
        # so comparing/indexing objects captured from two separate calls fails.
        paragraphs = doc.paragraphs
        target_i = self._find_heading_index(paragraphs, heading)
        if target_i is None:
            return False
        for q in paragraphs[target_i + 1:]:
            if q.style and q.style.name.startswith("Heading"):
                break
            if q.text.strip():
                original = q.text.strip()
                if len(original) <= 180:
                    return False
                shortened = self._shorten_text(prompt, original)
                q.clear()
                q.add_run(shortened)
                return True
        return False

    def _shorten_text(self, prompt, original):
        result = self.llm.json(
            "Rewrite the paragraph below to be noticeably shorter while preserving its actual "
            'meaning and facts — do not add new information. Return JSON only: {"text":"..."}.',
            f"REQUEST:\n{prompt}\nPARAGRAPH:\n{original}",
            None,
        )
        if isinstance(result, dict) and result.get("text"):
            return result["text"]
        # Generic fallback without an LLM: keep only the first sentence or two
        # of the actual paragraph rather than substituting unrelated text.
        sentences = re.split(r"(?<=[.!?])\s+", original)
        shortened = " ".join(sentences[:2]).strip()
        return shortened if shortened else original[:180].rstrip() + "…"

    def _detail_implementation_table(self, doc, table_index, prompt=""):
        if not (0 <= table_index < len(doc.tables)):
            return False
        table = doc.tables[table_index]
        if not table.rows or len(table.rows[0].cells) < 2:
            return False
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        details = self._generate_table_details(prompt, rows)
        table.add_column(Inches(2.6))
        table.cell(0, len(table.rows[0].cells) - 1).text = "Details"
        for r in range(1, len(table.rows)):
            activity = table.cell(r, 1).text.strip() if len(table.rows[r].cells) > 1 else table.cell(r, 0).text.strip()
            table.cell(r, len(table.rows[r].cells) - 1).text = details.get(
                str(r), f"Detailed execution and validation activities for {activity or 'this phase'}."
            )
        return True

    def _generate_table_details(self, prompt, rows):
        if len(rows) < 2:
            return {}
        result = self.llm.json(
            "For each data row (skip the header row) of the table below, write one detailed "
            "sentence describing that specific row, grounded only in its actual values. "
            'Return JSON only: {"details": {"<row_index_starting_at_1>":"...", ...}}.',
            f"REQUEST:\n{prompt}\nTABLE ROWS (row 0 is the header):\n{json.dumps(rows, ensure_ascii=False)}",
            {"details": {}},
        )
        if isinstance(result, dict) and isinstance(result.get("details"), dict):
            return {str(k): v for k, v in result["details"].items()}
        return {}

    def _replace_paragraph(self, doc, index, text):
        if 0 <= index < len(doc.paragraphs):
            p = doc.paragraphs[index]
            style = p.style
            p.clear()
            p.style = style
            p.add_run(text)
            return True
        return False

    def _delete_paragraph(self, doc, index):
        if 0 <= index < len(doc.paragraphs):
            p = doc.paragraphs[index]._element
            p.getparent().remove(p)
            return True
        return False

    def _insert_after_heading(self, doc, heading, paragraphs):
        p = self._find_heading_paragraph(doc, heading)
        if p is not None:
            anchor = p._p
            for text in paragraphs:
                np = doc.add_paragraph(text)
                anchor.addnext(np._p)
                anchor = np._p
            return bool(paragraphs)
        if paragraphs:
            # Heading not found in this document: append at the end rather
            # than silently failing the whole request.
            for text in paragraphs:
                doc.add_paragraph(text)
            return True
        return False

    def _replace_table(self, doc, table_index, rows):
        if not (0 <= table_index < len(doc.tables)):
            return False
        table = doc.tables[table_index]
        if not rows:
            return False
        for r, values in enumerate(rows):
            if r >= len(table.rows):
                table.add_row()
            for c, value in enumerate(values):
                if c < len(table.rows[r].cells):
                    table.rows[r].cells[c].text = str(value)
        return True

    def _find_numeric_column(self, rows):
        headers = rows[0]
        candidates = []
        for col, header in enumerate(headers):
            score = 0
            h = header.lower()
            if any(k in h for k in ["duration", "value", "amount", "count", "score", "week"]):
                score += 5
            values = []
            for row in rows[1:]:
                if col < len(row):
                    m = re.search(r"[-+]?\d+(?:\.\d+)?", row[col].replace(",", ""))
                    if m:
                        values.append(float(m.group()))
            if len(values) >= 2:
                candidates.append((score, col, values))
        return max(candidates, default=None, key=lambda x: x[0])

    def _add_chart_from_table(self, doc, table_index, chart_kind, title, replace_table=False):
        if not (0 <= table_index < len(doc.tables)):
            return False
        table = doc.tables[table_index]
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) < 2:
            return False
        numeric = self._find_numeric_column(rows)
        if not numeric:
            return False
        _, numeric_col, values = numeric
        labels = [row[1] if len(row) > 1 and row[1] else row[0] for row in rows[1:]]
        image_path = self._create_chart_image(labels, values, chart_kind, title or rows[0][numeric_col])
        if not image_path:
            return False
        parent = table._element.getparent()
        index = parent.index(table._element)
        if replace_table:
            backup = OxmlElement("w:p")
            parent.insert(index, backup)
            bp = Paragraph(backup, table._parent)
            run = bp.add_run(self.BACKUP_PREFIX + json.dumps(rows, ensure_ascii=False))
            run.font.hidden = True
            parent.remove(table._element)
            index += 1
        else:
            index += 1
        image_el = OxmlElement("w:p")
        parent.insert(index, image_el)
        ip = Paragraph(image_el, table._parent)
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.add_run().add_picture(image_path, width=Inches(6.1))
        cap_el = OxmlElement("w:p")
        parent.insert(index + 1, cap_el)
        cp = Paragraph(cap_el, table._parent)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run(f"{self.CHART_CAPTION_PREFIX}{title or 'Chart'}").font.hidden = True
        Path(image_path).unlink(missing_ok=True)
        return True

    def _restore_table_from_backup(self, doc):
        body = doc._body._element
        for i, child in enumerate(list(body)):
            text = "".join(child.itertext())
            if self.BACKUP_PREFIX not in text:
                continue
            raw = text.split(self.BACKUP_PREFIX, 1)[1]
            try:
                rows = json.loads(raw)
            except Exception:
                return False
            body.remove(child)
            remaining = list(body)
            for sibling in remaining[max(0, i - 1):i + 2]:
                sibling_text = "".join(sibling.itertext())
                if sibling.xpath('.//w:drawing') or self.CHART_CAPTION_PREFIX in sibling_text:
                    body.remove(sibling)
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    table.cell(r, c).text = str(value)
            return True
        return False

    def _remove_generated_visuals(self, doc):
        body = doc._body._element
        changed = False
        for child in list(body):
            text = "".join(child.itertext())
            if self.CHART_CAPTION_PREFIX in text:
                siblings = list(body)
                try:
                    idx = siblings.index(child)
                except ValueError:
                    continue
                body.remove(child)
                siblings = list(body)
                if idx > 0 and idx - 1 < len(siblings) and siblings[idx - 1].xpath('.//w:drawing'):
                    body.remove(siblings[idx - 1])
                changed = True
            elif self.BACKUP_PREFIX in text:
                body.remove(child)
                changed = True
        return changed

    def _create_chart_image(self, labels, values, kind, title):
        try:
            import matplotlib.pyplot as plt
            temp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            temp.close()
            path = temp.name
            fig, ax = plt.subplots(figsize=(8, 4.2))
            if kind == "pie":
                ax.pie(values, labels=labels, autopct="%1.0f%%")
            else:
                ax.bar(labels, values)
                ax.set_ylabel("Value")
                ax.tick_params(axis="x", rotation=25)
            ax.set_title(title)
            fig.tight_layout()
            fig.savefig(path, dpi=180, bbox_inches="tight")
            plt.close(fig)
            return path
        except Exception:
            return None

    def _polish_layout(self, doc):
        for section in doc.sections:
            section.top_margin = Inches(.7)
            section.bottom_margin = Inches(.7)
            section.left_margin = Inches(.8)
            section.right_margin = Inches(.8)
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and p.style.name == "Title":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif p.style and p.style.name.startswith("Heading"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
            elif p.style and p.style.name.startswith("List"):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
            elif not any(x in text for x in [self.CHART_CAPTION_PREFIX, self.BACKUP_PREFIX]):
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.08
        for table in doc.tables:
            table.autofit = True
            for row in table.rows:
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.space_after = Pt(2)
            if table.rows:
                for cell in table.rows[0].cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
