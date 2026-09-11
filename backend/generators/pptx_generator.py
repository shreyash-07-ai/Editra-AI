from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from ..llm import LLM


class PPTXGenerator:
    def __init__(self):
        self.llm = LLM()

    def generate(self, output_path, prompt, source_text="", template_path=None, existing=None, edit_plan=None):
        base = existing.get("path") if existing else template_path
        if base and Path(base).exists():
            shutil.copy2(base, output_path)
            if edit_plan:
                return self.apply_edit_plan(output_path, edit_plan)
            content = self._outline(prompt, source_text)
            prs = Presentation(output_path)
            self._append_content(prs, content)
            prs.save(output_path)
            return output_path

        prs = Presentation()
        content = self._outline(prompt, source_text)
        self._append_content(prs, content)
        prs.save(output_path)
        return output_path

    def _outline(self, prompt, source_text):
        return self.llm.json(
            """Create presentation slides grounded in the source material. Return JSON only:
{"slides":[{"title":"...","bullets":["..."],"visual":null|{"kind":"bar_chart"|"pie_chart"|"process","categories":[...],"values":[...],"series_name":"...","blocks":[...]}]}.
Do not invent unsupported numerical data. Use a chart only when the source contains suitable data.
Use a process/block visual for workflows or architectures when useful. If the user did not ask
for new slides, return {"slides":[]}.
""",
            f"REQUEST:\n{prompt}\nSOURCE:\n{source_text}",
            {"slides": []},
        )

    def _append_content(self, prs, content):
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        for s in content.get("slides", []):
            slide = prs.slides.add_slide(layout)
            if slide.shapes.title:
                slide.shapes.title.text = s.get("title", "")
            body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if body:
                body.text_frame.clear()
                for i, bullet in enumerate(s.get("bullets", [])):
                    p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
                    p.text = bullet
            if s.get("visual"):
                self._add_visual_to_slide(slide, s["visual"])

    def apply_edit_plan(self, path, plan):
        prs = Presentation(path)
        for op in plan.get("operations", []):
            kind = op.get("type")
            if kind == "replace_shape_text":
                self._replace_shape_text(prs, int(op.get("slide", 0)), op.get("old_text", ""), op.get("new_text", ""))
            elif kind == "replace_slide_text":
                self._replace_slide_text(prs, int(op.get("slide", 0)), op.get("text", ""))
            elif kind == "add_slide":
                self._add_slide(prs, op)
            elif kind == "delete_slide":
                self._delete_slide(prs, int(op.get("slide", 0)))
            elif kind == "add_visual":
                slide_no = int(op.get("slide", 0))
                if 1 <= slide_no <= len(prs.slides):
                    self._add_visual_to_slide(prs.slides[slide_no - 1], op.get("visual", {}))
        prs.save(path)
        return path

    def _replace_shape_text(self, prs, slide_no, old_text, new_text):
        if not (1 <= slide_no <= len(prs.slides)):
            return
        for shape in prs.slides[slide_no - 1].shapes:
            if hasattr(shape, "text") and shape.text == old_text:
                shape.text = new_text
                return

    def _replace_slide_text(self, prs, slide_no, text):
        if not (1 <= slide_no <= len(prs.slides)):
            return
        slide = prs.slides[slide_no - 1]
        text_shapes = [s for s in slide.shapes if hasattr(s, "text_frame")]
        if text_shapes:
            text_shapes[0].text_frame.text = text

    def _add_slide(self, prs, op):
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = op.get("title", "")
        body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
        if body:
            body.text_frame.clear()
            for i, bullet in enumerate(op.get("bullets", [])):
                p = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
                p.text = bullet
        if op.get("visual"):
            self._add_visual_to_slide(slide, op["visual"])

    def _delete_slide(self, prs, slide_no):
        if not (1 <= slide_no <= len(prs.slides)):
            return
        slide = prs.slides[slide_no - 1]
        slide._element.getparent().remove(slide._element)

    def _add_visual_to_slide(self, slide, visual):
        kind = visual.get("kind", "process")
        if kind in {"bar_chart", "pie_chart"}:
            categories = visual.get("categories", ["A", "B", "C"])
            values = visual.get("values", [1, 2, 3])
            if not categories or len(categories) != len(values):
                return
            data = CategoryChartData()
            data.categories = categories
            data.add_series(visual.get("series_name", "Value"), values)
            chart_type = XL_CHART_TYPE.PIE if kind == "pie_chart" else XL_CHART_TYPE.COLUMN_CLUSTERED
            slide.shapes.add_chart(chart_type, Inches(6.2), Inches(1.8), Inches(6.2), Inches(4.5), data)
            return

        blocks = visual.get("blocks", ["Input", "Analysis", "Output"])
        if not blocks:
            return
        y = 2.1
        for i, label in enumerate(blocks):
            x = 0.6 + i * min(3.4, 11.8 / max(1, len(blocks)))
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.6), Inches(1.0))
            shape.text_frame.text = str(label)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(16)
