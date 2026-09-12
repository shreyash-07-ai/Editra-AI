from pathlib import Path
import csv
import shutil

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font

from ..llm import LLM


class XLSXGenerator:
    """Edit an existing XLSX in-place on a copied version of the current artifact,
    or build a fresh workbook from source text / a CSV upload."""

    def __init__(self):
        self.llm = LLM()

    def generate(self, output_path, prompt, source_text="", existing=None, template_path=None, edit_plan=None):
        base = existing.get("path") if existing else template_path
        base_path = Path(base) if base else None

        if base_path and base_path.suffix.lower() == ".csv":
            # First touch on a CSV upload: normalize into an editable workbook,
            # then apply the requested edit as a plan (or a plain import if none).
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            self._import_csv(wb, base_path, "Sheet1")
            wb.save(output_path)
            if edit_plan:
                self.apply_edit_plan(output_path, edit_plan)
            return output_path

        if base_path and base_path.exists():
            shutil.copy2(base_path, output_path)
            if edit_plan:
                self.apply_edit_plan(output_path, edit_plan)
            else:
                addition = self._generate_addition(prompt, source_text)
                if addition:
                    wb = openpyxl.load_workbook(output_path)
                    self._apply_add_sheet(wb, addition)
                    wb.save(output_path)
            return output_path

        content = self.llm.json(
            "Create a spreadsheet from the source. Return JSON only: "
            '{"sheets":[{"name":"Sheet1","headers":["..."],"rows":[["..."]]}]}',
            f"REQUEST:\n{prompt}\nSOURCE:\n{source_text[:6000]}",
            {"sheets": [{"name": "Sheet1", "headers": ["Content"], "rows": [[line] for line in source_text.splitlines()[:50] if line.strip()]}]},
        )
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet in content.get("sheets", []) or [{"name": "Sheet1", "headers": [], "rows": []}]:
            self._write_sheet(wb, sheet.get("name", "Sheet1"), sheet.get("headers", []), sheet.get("rows", []))
        wb.save(output_path)
        return output_path

    def _import_csv(self, wb, csv_path, sheet_name):
        ws = wb.create_sheet(sheet_name)
        with open(csv_path, newline="", encoding="utf-8-sig", errors="replace") as f:
            for row in csv.reader(f):
                ws.append(row)
        self._style_header(ws)

    def _write_sheet(self, wb, name, headers, rows):
        ws = wb.create_sheet(name[:31] or "Sheet1")
        if headers:
            ws.append(headers)
        for row in rows:
            ws.append(row)
        self._style_header(ws)
        return ws

    def _style_header(self, ws):
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(40, max(10, length + 2))

    def _generate_addition(self, prompt, source_text):
        return self.llm.json(
            'Generate only the requested addition to an existing spreadsheet. '
            'Return JSON: {"name":"NewSheet","headers":["..."],"rows":[["..."]]}. '
            "Preserve existing sheets.",
            f"REQUEST:\n{prompt}\nSOURCE:\n{source_text[:4000]}",
            None,
        )

    def _apply_add_sheet(self, wb, addition):
        if not addition:
            return False
        self._write_sheet(wb, addition.get("name", "NewSheet"), addition.get("headers", []), addition.get("rows", []))
        return True

    def apply_edit_plan(self, path, plan):
        wb = openpyxl.load_workbook(path)
        operations = plan.get("operations", []) if isinstance(plan, dict) else []
        changed = False

        for op in operations:
            kind = op.get("type")
            if kind == "noop":
                continue
            elif kind == "set_cell":
                changed |= self._set_cell(wb, op.get("sheet"), op.get("cell", ""), op.get("value"))
            elif kind == "add_row":
                changed |= self._add_row(wb, op.get("sheet"), op.get("values", []))
            elif kind == "delete_row":
                changed |= self._delete_row(wb, op.get("sheet"), int(op.get("row", -1)))
            elif kind == "add_column":
                changed |= self._add_column(wb, op.get("sheet"), op.get("header", ""), op.get("values", []))
            elif kind == "add_sheet":
                changed |= self._apply_add_sheet(wb, op)
            elif kind == "rename_sheet":
                changed |= self._rename_sheet(wb, op.get("old_name", ""), op.get("new_name", ""))
            elif kind == "delete_sheet":
                changed |= self._delete_sheet(wb, op.get("name", ""))
            elif kind == "replace_range":
                changed |= self._replace_range(wb, op.get("sheet"), op.get("start_cell", "A1"), op.get("rows", []))

        if not changed and operations and not all(op.get("type") == "noop" for op in operations):
            raise ValueError("The requested spreadsheet operation could not be applied safely.")

        for ws in wb.worksheets:
            self._style_header(ws)
        wb.save(path)
        return path

    def _resolve_sheet(self, wb, name):
        if name and name in wb.sheetnames:
            return wb[name]
        return wb.worksheets[0] if wb.worksheets else None

    def _set_cell(self, wb, sheet, cell_ref, value):
        ws = self._resolve_sheet(wb, sheet)
        if not ws or not cell_ref:
            return False
        ws[cell_ref] = value
        return True

    def _add_row(self, wb, sheet, values):
        ws = self._resolve_sheet(wb, sheet)
        if not ws or not values:
            return False
        ws.append(values)
        return True

    def _delete_row(self, wb, sheet, row):
        ws = self._resolve_sheet(wb, sheet)
        if not ws or row < 1:
            return False
        ws.delete_rows(row)
        return True

    def _add_column(self, wb, sheet, header, values):
        ws = self._resolve_sheet(wb, sheet)
        if not ws:
            return False
        col = ws.max_column + 1
        if header:
            ws.cell(row=1, column=col, value=header)
        for i, value in enumerate(values, start=2):
            ws.cell(row=i, column=col, value=value)
        return True

    def _rename_sheet(self, wb, old_name, new_name):
        if old_name not in wb.sheetnames or not new_name:
            return False
        wb[old_name].title = new_name[:31]
        return True

    def _delete_sheet(self, wb, name):
        if name not in wb.sheetnames or len(wb.sheetnames) <= 1:
            return False
        del wb[name]
        return True

    def _replace_range(self, wb, sheet, start_cell, rows):
        ws = self._resolve_sheet(wb, sheet)
        if not ws or not rows:
            return False
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        col_letter, row_start = coordinate_from_string(start_cell)
        col_start = column_index_from_string(col_letter)
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                ws.cell(row=row_start + r, column=col_start + c, value=value)
        return True

    def export_csv(self, xlsx_path, csv_path, sheet_name=None):
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = self._resolve_sheet(wb, sheet_name)
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
        return csv_path
