from pathlib import Path
import csv as csv_module
import shutil
import subprocess
import tempfile

from docx import Document
from pptx import Presentation
import fitz
import openpyxl


def _ocr_image(path):
    try:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return ""


def _ocr_pdf_page(page):
    try:
        from PIL import Image
        import pytesseract
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(pix.tobytes("png"))
            tmp_path = tmp.name
        try:
            return pytesseract.image_to_string(Image.open(tmp_path))
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    except Exception:
        return ""


def extract_docx(path):
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            paragraphs.append({"index": i, "text": p.text, "style": p.style.name if p.style else ""})
    tables = []
    for ti, table in enumerate(doc.tables):
        tables.append({"index": ti, "rows": [[cell.text for cell in row.cells] for row in table.rows]})
    style = {}
    try:
        style["normal_font"] = doc.styles["Normal"].font.name
        style["normal_size"] = str(doc.styles["Normal"].font.size)
    except Exception:
        pass
    return {"type":"docx", "paragraphs":paragraphs, "tables":tables, "style":style,
            "paragraph_count":len(paragraphs), "table_count":len(tables)}


def extract_pptx(path):
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        items = []
        for si, shape in enumerate(slide.shapes):
            if hasattr(shape, "text") and shape.text.strip():
                items.append({"shape":si, "text":shape.text, "left":shape.left, "top":shape.top,
                              "width":shape.width, "height":shape.height, "shape_type":str(shape.shape_type)})
        slides.append({"number":i, "items":items, "layout":slide.slide_layout.name if slide.slide_layout else ""})
    return {"type":"pptx", "slide_count":len(prs.slides), "slides":slides}


def extract_pdf(path):
    doc = fitz.open(path)
    pages = []
    try:
        for i, page in enumerate(doc, 1):
            text = page.get_text("text")
            ocr_used = False
            if not text.strip():
                text = _ocr_pdf_page(page)
                ocr_used = bool(text.strip())
            pages.append({"number":i, "text":text, "ocr":ocr_used})
    finally:
        doc.close()
    return {"type":"pdf", "page_count":len(pages), "pages":pages}


def extract_image(path):
    text = _ocr_image(path)
    return {"type":"image", "path":str(path), "ocr_text":text, "ocr_available":bool(text.strip())}


MAX_ROWS_PREVIEW = 200


def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            rows.append(["" if v is None else v for v in row])
            if len(rows) >= MAX_ROWS_PREVIEW:
                break
        headers = rows[0] if rows else []
        sheets.append({
            "name": name,
            "headers": headers,
            "rows": rows[1:],
            "row_count": ws.max_row,
            "col_count": ws.max_column,
        })
    return {"type": "xlsx", "sheet_count": len(sheets), "sheets": sheets}


def extract_csv(path):
    rows = []
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        for row in csv_module.reader(f):
            rows.append(row)
            if len(rows) >= MAX_ROWS_PREVIEW:
                break
    headers = rows[0] if rows else []
    return {
        "type": "csv",
        "sheets": [{"name": "Sheet1", "headers": headers, "rows": rows[1:], "row_count": len(rows), "col_count": len(headers)}],
        "sheet_count": 1,
    }


def analyze(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx": return extract_docx(path)
    if ext == ".pptx": return extract_pptx(path)
    if ext == ".pdf": return extract_pdf(path)
    if ext in {".png", ".jpg", ".jpeg"}: return extract_image(path)
    if ext == ".xlsx": return extract_xlsx(path)
    if ext == ".csv": return extract_csv(path)
    if ext == ".ppt":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise ValueError("Legacy .ppt input requires LibreOffice for conversion to .pptx")
        out_dir = Path(path).parent / "converted"
        out_dir.mkdir(exist_ok=True)
        subprocess.run([soffice, "--headless", "--convert-to", "pptx", "--outdir", str(out_dir), str(path)], check=True)
        return extract_pptx(out_dir / (Path(path).stem + ".pptx"))
    raise ValueError(f"Unsupported file: {path}")
