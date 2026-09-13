from pathlib import Path
import json
import re
import uuid

from docx import Document
from docx.shared import Pt
import openpyxl

from .agents.rag_agent import RAGAgent
from .agents.research_agent import ResearchAgent
from .generators.docx_generator import DOCXGenerator
from .generators.pptx_generator import PPTXGenerator
from .generators.xlsx_generator import XLSXGenerator
from .llm import LLM
from .preview import preview
from .storage import MIME
from .config import OUTPUT_DIR


class QueryDrivenPipeline:
    """Source -> index -> retrieve -> analyze -> plan -> NEW editable artifact."""

    def __init__(self, store):
        self.store = store
        self.llm = LLM()
        self.rag = RAGAgent()
        self.research = ResearchAgent()
        self.doc_gen = DOCXGenerator()
        self.ppt_gen = PPTXGenerator()
        self.xlsx_gen = XLSXGenerator()

    @staticmethod
    def _text(value):
        return str(value).strip() if value is not None else ""

    def _chunks_from_analysis(self, analysis, source_name):
        chunks = []
        kind = analysis.get("type")

        def add(text, section="", page=0, content_type="text", metadata=None):
            text = self._text(text)
            if not text:
                return
            chunks.append({
                "content": text,
                "section": section,
                "page_number": page,
                "content_type": content_type,
                "source": source_name,
                "metadata": metadata or {},
            })

        if kind == "pdf":
            for page in analysis.get("pages", []):
                page_no = int(page.get("number", 0))
                text = self._text(page.get("text"))
                if text:
                    self._split_and_add(add, text, page_no, source_name)
                for table in page.get("tables", []):
                    rows = table.get("rows", [])
                    if rows:
                        add("TABLE:\n" + json.dumps(rows, ensure_ascii=False), "Table", page_no, "table", {"table_index": table.get("index", 0)})
                if page.get("visual_text"):
                    add(page["visual_text"], "Visual", page_no, "chart", {})
        elif kind == "docx":
            for p in analysis.get("paragraphs", []):
                section = p.get("text", "") if p.get("style", "").startswith("Heading") else ""
                add(p.get("text"), section, 0, "heading" if section else "text", {"paragraph_index": p.get("index", 0)})
            for table in analysis.get("tables", []):
                add("TABLE:\n" + json.dumps(table.get("rows", []), ensure_ascii=False), "Table", 0, "table", {"table_index": table.get("index", 0)})
        elif kind == "pptx":
            for slide in analysis.get("slides", []):
                text = " | ".join(self._text(x.get("text")) for x in slide.get("items", []) if self._text(x.get("text")))
                add(text, f"Slide {slide.get('number')}", int(slide.get("number", 0)), "slide", {})
        elif kind in ("xlsx", "csv"):
            for sheet in analysis.get("sheets", []):
                rows = [sheet.get("headers", [])] + sheet.get("rows", [])
                add("SHEET " + self._text(sheet.get("name")) + "\n" + json.dumps(rows[:200], ensure_ascii=False), self._text(sheet.get("name")), 0, "table", {"row_count": sheet.get("row_count", len(rows))})
        elif kind == "image":
            add(analysis.get("ocr_text"), "Image", 0, "image", {})
        return chunks

    def _split_and_add(self, add_fn, text, page, source_name):
        # Preserve page boundaries and use overlap so retrieval keeps context.
        size, overlap = 1200, 180
        text = re.sub(r"\s+", " ", text).strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                boundary = text.rfind(". ", start, end)
                if boundary > start + 500:
                    end = boundary + 1
            add_fn(text[start:end], "", page, "text", {"source_name": source_name})
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)

    def _query_analysis(self, query, description):
        fallback = {
            "output_type": self._output_type(query),
            "needs_web": any(x in query.lower() for x in ["current", "latest", "web", "today", "recent", "trends", "research"]),
            "intent": query,
            "focus": query,
            "keywords": [],
        }
        system = """You are a query-understanding agent for a document RAG system. Return JSON only:
{"output_type":"docx|pptx|xlsx","needs_web":true|false,"intent":"...","focus":"...","keywords":["..."]}.
Choose pptx for presentations/slide decks, xlsx for spreadsheets/tables/numerical work, otherwise docx.
needs_web is true only when the user asks for current/external information or research beyond the uploaded source.
The uploaded document is source material, not an output template."""
        return self.llm.json(system, f"DOCUMENT DESCRIPTION:\n{description}\nQUERY:\n{query}", fallback)

    @staticmethod
    def _output_type(query):
        p = query.lower()
        if any(x in p for x in ["presentation", "powerpoint", "pptx", "slide deck", "slides"]):
            return "pptx"
        if any(x in p for x in ["spreadsheet", "excel", "xlsx", "table of data", "numerical analysis"]):
            return "xlsx"
        return "docx"

    def _semantic_analysis(self, query, description, retrieved, web_results):
        context = "\n\n".join(
            f"[DOC {i+1}] {x.get('source','')} p.{x.get('page_number',0)}\n{x.get('content','')}"
            for i, x in enumerate(retrieved)
        )
        web = "\n\n".join(
            f"[WEB] {x.get('title','')}\n{x.get('content','')}\nURL: {x.get('url','')}"
            for x in web_results
        )
        fallback = {
            "title": "Query Analysis",
            "summary": "Analysis generated from the retrieved source context.",
            "findings": [x.get("content", "")[:500] for x in retrieved[:5]],
            "risks": [], "recommendations": [], "gaps": [], "numbers": [],
        }
        system = """Perform semantic analysis before document generation. Return JSON only with keys:
title, summary, findings[], risks[], recommendations[], gaps[], numbers[].
Use only retrieved document context and supplied web results. Do not invent facts or numbers.
Classify direct answers, supporting evidence, trends, risks, recommendations, contradictions and missing information.
If a web result exists, keep web-derived claims distinct from document-derived claims."""
        user = f"QUERY:\n{query}\nDOCUMENT DESCRIPTION:\n{description}\nRETRIEVED CONTEXT:\n{context}\nWEB CONTEXT:\n{web}"
        return self.llm.json(system, user, fallback)

    def _generation_prompt(self, query, description, analysis, retrieved, web_results):
        doc_context = "\n\n".join(
            f"SOURCE={x.get('source','')} PAGE={x.get('page_number',0)} TYPE={x.get('content_type','text')}\n{x.get('content','')}"
            for x in retrieved
        )
        web_context = "\n\n".join(
            f"WEB={x.get('title','')} URL={x.get('url','')}\n{x.get('content','')}" for x in web_results
        )
        return f"""Create a NEW editable artifact for this query. NEVER copy, return, or reproduce the uploaded source document.
The query controls the structure and scope. Use only the retrieved context as the primary document evidence.
Document description: {description}
Query: {query}
Semantic analysis: {json.dumps(analysis, ensure_ascii=False)}
Retrieved document context:\n{doc_context}
Web context (only if present):\n{web_context}
Clearly distinguish document evidence from web evidence. Preserve source numbers exactly and say when something is unavailable.
"""

    def _generate_docx(self, output_path, prompt):
        content = self.llm.json(
            """Generate a completely new professional report. Return JSON only:
{"title":"...","sections":[{"heading":"...","paragraphs":["..."],"bullets":["..."]}],"sources":["..."]}.
The structure must answer the user's query rather than mirror the source document. Do not invent facts or numbers.""",
            prompt,
            {"title": "Query-Specific Analysis", "sections": [{"heading": "Findings", "paragraphs": ["The retrieved context did not produce a structured LLM response."]}], "sources": []},
        )
        doc = Document()
        doc.add_heading(content.get("title", "Query-Specific Analysis"), 0)
        for section in content.get("sections", []):
            doc.add_heading(section.get("heading", "Section"), 1)
            for p in section.get("paragraphs", []):
                doc.add_paragraph(str(p))
            for b in section.get("bullets", []):
                doc.add_paragraph(str(b), style="List Bullet")
        if content.get("sources"):
            doc.add_heading("Sources", 1)
            for source in content["sources"]:
                doc.add_paragraph(str(source))
        doc.save(output_path)

    def _add_xlsx_sources(self, path, sources):
        if not sources:
            return
        wb = openpyxl.load_workbook(path)
        ws = wb.create_sheet("Sources")
        ws.append(["Source"])
        for source in sources:
            ws.append([source])
        wb.save(path)

    def run(self, description, query, upload_paths, output_type="Auto"):
        if not upload_paths:
            return {"message": "Upload a source document before generating a query-specific artifact."}
        if not query.strip():
            return {"message": "Enter a query describing what the new document should contain."}

        cid = uuid.uuid4().hex[:10]
        document_id = uuid.uuid5(uuid.NAMESPACE_URL, "editra:" + cid).hex
        analyses, chunks = [], []
        for raw in upload_paths:
            path = Path(raw)
            try:
                from .processors import analyze
                analysis = analyze(path)
                analyses.append(analysis)
                chunks.extend(self._chunks_from_analysis(analysis, path.name))
            except Exception as exc:
                return {"message": f"[EXTRACTION] Failed for {path.name}: {exc}"}

        print(f"[INGESTION] Document received: {len(upload_paths)} file(s)")
        print(f"[EXTRACTION] {sum(a.get('page_count', a.get('paragraph_count', a.get('slide_count', 1))) for a in analyses)} pages/units processed")
        print(f"[CHUNKING] {len(chunks)} chunks created")
        try:
            index_stats = self.rag.index_document(document_id, chunks)
        except Exception as exc:
            return {"message": f"[EMBEDDING/VECTOR DB] Indexing failed: {exc}"}
        print(f"[EMBEDDING] {index_stats['chunks']} vectors generated")
        print(f"[VECTOR DB] {index_stats['stored']} vectors stored" if index_stats["enabled"] else "[VECTOR DB] Pinecone not configured; semantic local fallback used")

        qinfo = self._query_analysis(query, description)
        chosen_type = self._output_type(query) if output_type == "Auto" else output_type.lower()
        if chosen_type not in {"docx", "pptx", "xlsx"}:
            chosen_type = "docx"
        print(f"[QUERY] Query received: {query}")
        print(f"[QUERY ANALYSIS] Intent={qinfo.get('intent')} output={chosen_type}")

        retrieved = self.rag.retrieve(query, document_id=document_id, top_k=8)
        print(f"[RETRIEVAL] {len(retrieved)} relevant chunks retrieved")

        web_results = []
        if qinfo.get("needs_web"):
            rr = self.research.search(query)
            web_results = rr.get("results", [])
        print(f"[WEB] Search required: {'YES' if qinfo.get('needs_web') else 'NO'}")
        print(f"[WEB] {len(web_results)} sources retrieved")

        analysis = self._semantic_analysis(query, description, retrieved, web_results)
        print("[ANALYSIS] Context analyzed")
        generation_prompt = self._generation_prompt(query, description, analysis, retrieved, web_results)
        aid, output_path = self.store.new_output({"docx": ".docx", "pptx": ".pptx", "xlsx": ".xlsx"}[chosen_type], 1)
        try:
            if chosen_type == "docx":
                self._generate_docx(output_path, generation_prompt)
            elif chosen_type == "pptx":
                self.ppt_gen.generate(output_path, query, generation_prompt)
            else:
                self.xlsx_gen.generate(output_path, query, generation_prompt)
                self._add_xlsx_sources(output_path, [x.get("url") for x in web_results if x.get("url")])
        except Exception as exc:
            return {"message": f"[GENERATION] Failed: {exc}"}

        ext = output_path.suffix.lower()
        try:
            from .agents.validation_agent import ValidationAgent
            ok, checks = ValidationAgent().validate(output_path, chosen_type)
        except Exception as exc:
            return {"message": f"[VALIDATION] Failed: {exc}"}
        if not ok:
            return {"message": "Generated output failed validation: " + "; ".join(checks)}

        pdir = self.store.preview_dir(aid)
        images = preview(output_path, chosen_type, pdir) if chosen_type in {"docx", "pptx"} else []
        sources = [x.get("source", "") for x in retrieved if x.get("source")]
        sources.extend(x.get("url", "") for x in web_results if x.get("url"))
        artifact = {
            "id": aid, "conversation_id": cid, "filename": output_path.name, "path": str(output_path),
            "artifact_type": chosen_type, "version": 1, "mime": MIME.get(ext, "application/octet-stream"),
            "preview_type": "images" if images else "text",
            "preview": images if images else f"New query-specific {chosen_type.upper()} generated from {len(retrieved)} retrieved chunks.",
            "structure": {"query": query, "description": description, "analysis": analysis, "retrieved_chunks": len(retrieved), "web_sources": len(web_results)},
            "sources": list(dict.fromkeys(sources)),
        }
        self.store.register(artifact)
        print(f"[GENERATION] Output structure created")
        print(f"[GENERATION] Editable {chosen_type.upper()} generated")
        return {"message": f"Done — created a NEW query-specific **{artifact['filename']}**. The uploaded source was used only as reference material.", "artifact": artifact, "sources": artifact["sources"]}
