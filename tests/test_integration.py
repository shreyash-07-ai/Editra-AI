"""End-to-end smoke tests for the iterative editing workflow.

These run without a GEMINI_API_KEY: the LLM wrapper returns each call's
fallback value, so structural plumbing (routing, versioning, file
generation, validation, undo, restore-original, format export) is fully
exercised even though free-text content generation degrades to the
fallback content.
"""
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from pptx import Presentation
import openpyxl
import fitz

from backend.orchestrator import EditraOrchestrator
from backend.storage import ArtifactStore


def make_orch():
    store = ArtifactStore()
    return EditraOrchestrator(store), store


def make_sample_docx(path):
    doc = Document()
    doc.add_heading("Sample Report", 0)
    doc.add_heading("1. Introduction", 1)
    doc.add_paragraph("This is a sample report used for testing.")
    doc.add_heading("3. Proposed Solution", 1)
    doc.add_paragraph("The proposed solution automates document generation and editing.")
    doc.add_heading("5. Implementation Plan", 1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Phase"
    table.rows[0].cells[1].text = "Activity"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "Kickoff"
    doc.add_heading("6. Success Metrics", 1)
    doc.add_paragraph(
        "We will measure success using a broad and deliberately long-winded set of metrics "
        "that cover adoption, engagement, retention, satisfaction, and overall business impact "
        "across every team involved in the rollout, tracked on a rolling monthly basis."
    )
    doc.save(path)


def make_sample_pptx(path):
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Intro Slide"
    prs.save(path)


def make_sample_pdf(path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from a sample PDF report about quarterly performance.")
    doc.save(path)
    doc.close()


def make_sample_csv(path):
    Path(path).write_text("Name,Score\nAlice,90\nBob,85\n", encoding="utf-8")


def test_pdf_to_docx():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        pdf_path = Path(d) / "input.pdf"
        make_sample_pdf(pdf_path)
        result = orch.run("Turn this into a Word document.", [str(pdf_path)], None, [])
        assert result.get("artifact"), result.get("message")
        assert result["artifact"]["artifact_type"] == "docx"
        assert Path(result["artifact"]["path"]).exists()


def test_docx_edit_in_place():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        r1 = orch.run("Add an executive summary at the beginning.", [str(docx_path)], None, [])
        assert r1.get("artifact"), r1.get("message")
        assert r1["artifact"]["version"] == 1
        r2 = orch.run("Shorten the success metrics section.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        assert r2["artifact"]["version"] == 2


def test_pdf_to_pptx():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        pdf_path = Path(d) / "input.pdf"
        make_sample_pdf(pdf_path)
        result = orch.run("Convert this report into a presentation.", [str(pdf_path)], None, [])
        assert result.get("artifact"), result.get("message")
        assert result["artifact"]["artifact_type"] == "pptx"


def test_pptx_edit_in_place():
    orch, store = make_orch()
    # PPTX has no hardcoded offline patterns (unlike a couple of DOCX ops), so
    # exercising the full edit-in-place pipeline here means stubbing the LLM
    # boundary with a deterministic plan rather than requiring a live API key.
    with TemporaryDirectory() as d:
        pptx_path = Path(d) / "input.pptx"
        make_sample_pptx(pptx_path)
        orch.editing.llm.json = lambda system, user, fallback: {
            "operations": [{"type": "replace_shape_text", "slide": 1, "old_text": "Intro Slide", "new_text": "Introduction"}]
        }
        r1 = orch.run("Make the title more professional.", [str(pptx_path)], None, [])
        assert r1.get("artifact"), r1.get("message")
        orch.editing.llm.json = lambda system, user, fallback: {
            "operations": [{"type": "add_slide", "after_slide": 1, "title": "Conclusion", "bullets": ["Thank you"]}]
        }
        r2 = orch.run("Add a conclusion slide.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        assert r2["artifact"]["version"] == 2


def test_csv_to_modified_xlsx():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        csv_path = Path(d) / "input.csv"
        make_sample_csv(csv_path)
        orch.editing.llm.json = lambda system, user, fallback: {
            "operations": [{"type": "add_sheet", "name": "Summary", "headers": ["Metric", "Value"], "rows": [["Average", 87.5]]}]
        }
        r1 = orch.run("Add a new sheet summarizing the scores.", [str(csv_path)], None, [])
        assert r1.get("artifact"), r1.get("message")
        assert r1["artifact"]["artifact_type"] == "xlsx"
        wb = openpyxl.load_workbook(r1["artifact"]["path"])
        assert "Sheet1" in wb.sheetnames
        orch.editing.llm.json = lambda system, user, fallback: {
            "operations": [{"type": "add_row", "sheet": "Sheet1", "values": ["Total", 175]}]
        }
        r2 = orch.run("Add a total row.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        assert r2["artifact"]["version"] == 2


def test_version_chain_v1_v2_v3():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        r1 = orch.run("Add an executive summary.", [str(docx_path)], None, [])
        r2 = orch.run("Make the implementation plan more detailed.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        r3 = orch.run("Make the success metrics section shorter.", [], r2["artifact"], [])
        assert r3.get("artifact"), r3.get("message")
        assert [r["artifact"]["version"] for r in (r1, r2, r3)] == [1, 2, 3]


def test_current_version_to_pdf():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        r1 = orch.run("Add an executive summary.", [str(docx_path)], None, [])
        r2 = orch.run("Give me a PDF version of this.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        assert r2["artifact"]["artifact_type"] == "pdf"
        assert r2["artifact"]["editable_source"]["id"] == r1["artifact"]["id"]
        # Editing after a PDF export should resolve back to the editable DOCX.
        r3 = orch.run("Shorten the success metrics section.", [], r2["artifact"], [])
        assert r3.get("artifact"), r3.get("message")
        assert r3["artifact"]["artifact_type"] == "docx"


def test_undo_and_restore_original():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        r1 = orch.run("Add an executive summary.", [str(docx_path)], None, [])
        r2 = orch.run("Make the implementation plan more detailed.", [], r1["artifact"], [])
        assert r2.get("artifact"), r2.get("message")
        undo = orch.run("Undo the last change.", [], r2["artifact"], [])
        assert undo.get("artifact")
        assert undo["artifact"]["id"] == r1["artifact"]["id"]
        restore = orch.run("Go back to the original.", [], r2["artifact"], [])
        assert restore.get("artifact"), restore.get("message")
        assert restore["artifact"]["version"] == r2["artifact"]["version"] + 1


def test_invalid_corrupted_input_fails_gracefully():
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        bad_path = Path(d) / "broken.docx"
        bad_path.write_bytes(b"not a real docx file")
        result = orch.run("Fix this document.", [str(bad_path)], None, [])
        # Must not raise; the working document (none yet) is untouched either way.
        assert "message" in result


def test_create_ppt_from_plain_text_prompt_no_trailing_words():
    # Regression test: "ppt" as the very last word (no trailing space/period)
    # must still be recognized as a PPT request.
    orch, store = make_orch()
    result = orch.run("Create a PPT", [], None, [])
    assert result.get("artifact"), result.get("message")
    assert result["artifact"]["artifact_type"] == "pptx"


def test_upload_docx_request_ppt_on_first_message():
    # Regression test: uploading a DOCX and asking for a PPT on the very first
    # message must produce a PPTX, not silently edit the DOCX in its own format.
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        result = orch.run("Convert this into a PowerPoint presentation.", [str(docx_path)], None, [])
        assert result.get("artifact"), result.get("message")
        assert result["artifact"]["artifact_type"] == "pptx"


def test_unmatched_edit_request_reports_honestly_instead_of_false_success():
    # Regression test: when no operation can be safely determined (e.g. no
    # GEMINI_API_KEY and no hardcoded pattern matches), the app must say so
    # instead of returning a "Done!" message with an unmodified file.
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        result = orch.run(
            "Please rephrase paragraph fourteen to sound more energetic.",
            [str(docx_path)], None, [],
        )
        assert not result.get("artifact")
        assert "message" in result and result["message"]


def test_executive_summary_is_content_grounded_not_boilerplate():
    # Regression test: the executive summary must not hardcode unrelated
    # fictional company/product text regardless of the actual document.
    orch, store = make_orch()
    with TemporaryDirectory() as d:
        docx_path = Path(d) / "input.docx"
        make_sample_docx(docx_path)
        result = orch.run("Add an executive summary.", [str(docx_path)], None, [])
        assert result.get("artifact"), result.get("message")
        doc = Document(result["artifact"]["path"])
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "NovaEdge" not in full_text
        assert "Sample Report" in full_text
